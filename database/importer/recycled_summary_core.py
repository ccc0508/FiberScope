from __future__ import annotations

import hashlib
import json
from collections import defaultdict
from dataclasses import asdict, dataclass
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable
from zipfile import ZipFile
from xml.etree import ElementTree

from openpyxl import load_workbook


FILE_NAME = "数据汇总(1).xlsx"
SHEET_NAME = "Sheet1"
SOURCE_VERSION = "正式汇总版"
DATASET_CODE = "RECYCLED_SUMMARY_2025"
STUDY_NAME = "再生纸浆循环回用实验"
EXPECTED_SHA256 = "1299ff1cb43d79e8bff21b94b28ca47e6cd281a7e3e6c3fdaed699310a217e97"

EXPECTED_HEADERS = (
    "循环次数",
    "0.2-0.3mm纤维占比（%）",
    "0.3-0.45（%）",
    "0.45-0.67（%）",
    "0.67-1（%）",
    "1-1.5（%）",
    "1.5-2.24（%）",
    "2.24-3.35（%）",
    "3.35-5.01（%）",
    "5.01-7.5（%）",
    "＞7.5（%）",
    "算数平均长/mm",
    "长度加权长/mm",
    "宽度/μm",
    "粗度/(mg/m)",
    "扭结角度/度",
    "扭结纤维占比/%",
    "弯曲/%",
    "分丝帚化长度比/%",
    "折断纤维比/%",
    "细小纤维长度比/%",
    "细小纤维加权长度比/%",
    "细小纤维面积比/%",
    "抗张强度（kN/m）",
    "耐破强度（kPa）",
    "撕裂强度（mN）",
)

VALUE_COLUMNS = (
    "length_0_2_0_3_pct",
    "length_0_3_0_45_pct",
    "length_0_45_0_67_pct",
    "length_0_67_1_pct",
    "length_1_1_5_pct",
    "length_1_5_2_24_pct",
    "length_2_24_3_35_pct",
    "length_3_35_5_01_pct",
    "length_5_01_7_5_pct",
    "length_gt_7_5_pct",
    "fiber_length_arithmetic_mm",
    "fiber_length_weighted_mm",
    "fiber_width_um",
    "fiber_coarseness_mg_m",
    "kink_angle_deg",
    "kinked_fiber_pct",
    "curved_fiber_pct",
    "fibrillation_length_pct",
    "broken_fiber_pct",
    "fines_length_pct",
    "fines_weighted_length_pct",
    "fines_area_pct",
    "tensile_strength_kn_m",
    "burst_strength_kpa",
    "tear_strength_mn",
)

BLOCK_SPECS = (
    ("G01", 1, 1, None, "美废12#", "同项目实验说明：以进口美废12#再生纸浆为原料", 3, 5, 2, 16),
    ("G02", 2, 1, None, "美废12#", "同项目实验说明：以进口美废12#再生纸浆为原料", 3, 5, 17, 31),
    ("G03", 3, 1, None, "美废12#", "同项目实验说明：以进口美废12#再生纸浆为原料", 3, 5, 32, 46),
    ("G04", 4, 1, None, "美废12#", "同项目实验说明：以进口美废12#再生纸浆为原料", 3, 5, 47, 61),
    ("G05", 5, 1, None, "美废12#", "同项目实验说明：以进口美废12#再生纸浆为原料", 3, 5, 62, 76),
    ("G06", 6, 1, None, "美废12#", "同项目实验说明：以进口美废12#再生纸浆为原料", 3, 3, 77, 85),
    ("G07", 7, 1, None, "美废12#", "同项目实验说明：以进口美废12#再生纸浆为原料", 3, 3, 86, 94),
    ("G08", 8, 1, None, "美废12#", "同项目实验说明：以进口美废12#再生纸浆为原料", 3, 3, 95, 103),
    ("G09", 9, 1, None, "美废12#", "同项目实验说明：以进口美废12#再生纸浆为原料", 3, 3, 104, 112),
    ("G10", 10, 2, "陈莉", "美废12#", "工作簿第134行批注：原料为美废12#", 4, 1, 114, 117),
    ("G11", 11, 2, "陈莉", "美废13#", "工作簿第135行批注及第118-121行黄色填充", 4, 1, 118, 121),
    ("G12", 12, 2, "陈莉", "美废12#", "工作簿第134行批注；非黄色数据组", 4, 1, 122, 125),
    ("G13", 13, 2, "陈莉", "美废12#", "工作簿第134行批注；非黄色数据组", 4, 1, 126, 129),
    ("G14", 14, 2, "陈莉", "美废12#", "工作簿第134行批注；非黄色数据组", 4, 1, 130, 133),
)


class RecycledImportError(ValueError):
    pass


@dataclass(frozen=True)
class DataBlock:
    block_code: str
    block_name: str
    block_order: int
    section_no: int
    contributor_name: str | None
    raw_material_grade: str | None
    material_assignment_basis: str
    condition_label: str | None
    metadata_status: str
    cycle_count: int
    expected_replicates_per_cycle: int
    source_start_row: int
    source_end_row: int
    notes: str


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _core_properties(path: Path) -> tuple[str | None, str | None]:
    ns = {
        "dc": "http://purl.org/dc/elements/1.1/",
        "cp": "http://schemas.openxmlformats.org/package/2006/metadata/core-properties",
    }
    with ZipFile(path) as archive:
        root = ElementTree.fromstring(archive.read("docProps/core.xml"))
    creator = root.findtext("dc:creator", namespaces=ns)
    modifier = root.findtext("cp:lastModifiedBy", namespaces=ns)
    return creator, modifier


def _decimal(value: Any, row: int, column: int) -> Decimal:
    if value is None or isinstance(value, bool):
        raise RecycledImportError(f"Sheet1!{row},{column}必须为数值，当前为空或布尔值")
    try:
        return Decimal(str(value))
    except (InvalidOperation, ValueError, TypeError) as exc:
        raise RecycledImportError(
            f"Sheet1!{row},{column}必须为数值，当前为{value!r}"
        ) from exc


def _yellow_fill(cell: Any) -> bool:
    fill = cell.fill
    return (
        fill.fill_type == "solid"
        and fill.fgColor.type == "rgb"
        and str(fill.fgColor.rgb).upper() in {"FFFFFF00", "00FFFF00"}
    )


def _validate_value_ranges(row: int, values: list[Decimal]) -> None:
    percentages = values[:10] + values[15:22]
    if any(value < 0 or value > 100 for value in percentages):
        raise RecycledImportError(f"第{row}行存在不在0-100范围内的百分比")
    if values[14] < 0 or values[14] > 180:
        raise RecycledImportError(f"第{row}行扭结角度不在0-180度范围内")
    if any(value < 0 for value in values[10:15] + values[22:25]):
        raise RecycledImportError(f"第{row}行存在负的纤维或纸张性能值")


def validate_and_build(data_dir: Path) -> dict[str, Any]:
    path = data_dir / FILE_NAME
    if not path.is_file():
        raise RecycledImportError(f"找不到来源文件：{path}")

    actual_sha256 = _sha256(path)
    if actual_sha256 != EXPECTED_SHA256:
        raise RecycledImportError(
            "来源文件SHA-256与已审计版本不一致："
            f"期望{EXPECTED_SHA256}，实际{actual_sha256}"
        )

    workbook = load_workbook(path, data_only=False, read_only=False)
    if tuple(workbook.sheetnames) != ("Sheet1", "Sheet2", "Sheet3"):
        raise RecycledImportError(f"工作表结构变化：{workbook.sheetnames!r}")
    worksheet = workbook[SHEET_NAME]

    headers = tuple(worksheet.cell(1, column).value for column in range(1, 27))
    if headers != EXPECTED_HEADERS:
        raise RecycledImportError("Sheet1表头与已审计版本不一致")

    formula_cells = [
        cell.coordinate
        for row in worksheet.iter_rows()
        for cell in row
        if cell.data_type == "f"
    ]
    if formula_cells:
        raise RecycledImportError(f"来源表不应包含公式：{formula_cells[:10]!r}")

    if worksheet["A113"].value != "陈莉":
        raise RecycledImportError("第113行分段标记应为“陈莉”")
    if worksheet["B134"].value != "原料：美废12#":
        raise RecycledImportError("第134行原料批注与已审计版本不一致")
    if worksheet["B135"].value != "标黄是美废13#":
        raise RecycledImportError("第135行黄色含义批注与已审计版本不一致")

    yellow_rows = [row for row in range(2, 134) if _yellow_fill(worksheet.cell(row, 1))]
    if yellow_rows != [118, 119, 120, 121]:
        raise RecycledImportError(f"黄色数据行应为118-121，当前为{yellow_rows!r}")

    blocks: list[DataBlock] = []
    observations: list[dict[str, Any]] = []
    quality_rows: list[dict[str, Any]] = []
    all_source_rows: list[int] = []

    for spec in BLOCK_SPECS:
        (
            block_code,
            block_order,
            section_no,
            contributor_name,
            raw_material_grade,
            material_assignment_basis,
            cycle_count,
            expected_replicates,
            start_row,
            end_row,
        ) = spec
        block = DataBlock(
            block_code=block_code,
            block_name=f"数据块{block_order:02d}",
            block_order=block_order,
            section_no=section_no,
            contributor_name=contributor_name,
            raw_material_grade=raw_material_grade,
            material_assignment_basis=material_assignment_basis,
            condition_label=None,
            metadata_status="unknown",
            cycle_count=cycle_count,
            expected_replicates_per_cycle=expected_replicates,
            source_start_row=start_row,
            source_end_row=end_row,
            notes="源表未提供正式实验条件；仅按连续数据块和源表位置建立稳定编号",
        )
        blocks.append(block)

        replicate_numbers: defaultdict[int, int] = defaultdict(int)
        for row in range(start_row, end_row + 1):
            cycle_value = _decimal(worksheet.cell(row, 1).value, row, 1)
            if cycle_value != cycle_value.to_integral_value():
                raise RecycledImportError(f"第{row}行循环次数必须为整数")
            cycle_no = int(cycle_value)
            if cycle_no < 1 or cycle_no > cycle_count:
                raise RecycledImportError(
                    f"第{row}行循环次数{cycle_no}超出数据块{block_code}预期范围"
                )

            replicate_numbers[cycle_no] += 1
            values = [_decimal(worksheet.cell(row, column).value, row, column) for column in range(2, 27)]
            _validate_value_ranges(row, values)
            observation = {
                "block_code": block_code,
                "cycle_no": cycle_no,
                "replicate_no": replicate_numbers[cycle_no],
                "source_row": row,
            }
            observation.update(dict(zip(VALUE_COLUMNS, values, strict=True)))
            observations.append(observation)
            all_source_rows.append(row)

            distribution_sum = sum(values[:10], Decimal("0"))
            if abs(distribution_sum - Decimal("100")) > Decimal("1"):
                quality_rows.append(
                    {
                        "source_row": row,
                        "block_code": block_code,
                        "cycle_no": cycle_no,
                        "replicate_no": replicate_numbers[cycle_no],
                        "distribution_sum_pct": distribution_sum,
                        "deviation_pct": distribution_sum - Decimal("100"),
                    }
                )

        expected_cycles = set(range(1, cycle_count + 1))
        if set(replicate_numbers) != expected_cycles:
            raise RecycledImportError(f"数据块{block_code}的循环次数不完整")
        if any(count != expected_replicates for count in replicate_numbers.values()):
            raise RecycledImportError(f"数据块{block_code}的重复数与已审计结构不一致")

    expected_rows = list(range(2, 113)) + list(range(114, 134))
    if all_source_rows != expected_rows:
        raise RecycledImportError("数值行范围与已审计版本不一致")
    if len(blocks) != 14 or len(observations) != 131 or len(quality_rows) != 4:
        raise RecycledImportError("数据块、观测或质量复核数量与已审计版本不一致")

    creator, modifier = _core_properties(path)
    return {
        "source_path": path,
        "dataset": {
            "dataset_code": DATASET_CODE,
            "study_name": STUDY_NAME,
            "file_name": FILE_NAME,
            "sheet_name": SHEET_NAME,
            "source_version": SOURCE_VERSION,
            "file_sha256": actual_sha256,
            "workbook_creator": creator,
            "workbook_last_modified_by": modifier,
            "notes": (
                "来源表无样品ID和正式实验条件；按连续数据块生成G01-G14和数据块01-14。"
                "原始数值逐行保留，长度分布不自动归一化。"
            ),
        },
        "blocks": [asdict(block) for block in blocks],
        "observations": observations,
        "quality_rows": quality_rows,
    }


def _json_value(value: Any) -> Any:
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, Path):
        return str(value)
    return value


def build_report(dataset: dict[str, Any]) -> dict[str, Any]:
    blocks = dataset["blocks"]
    observations = dataset["observations"]
    summary_points = len({(row["block_code"], row["cycle_no"]) for row in observations})
    material_counts: defaultdict[str, int] = defaultdict(int)
    for row in blocks:
        material_counts[row["raw_material_grade"] or "未确认"] += 1
    return {
        "status": "PASS",
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "source": dataset["dataset"],
        "counts": {
            "datasets": 1,
            "blocks": len(blocks),
            "observations": len(observations),
            "summary_points": summary_points,
            "quality_review_rows": len(dataset["quality_rows"]),
        },
        "block_material_counts": dict(material_counts),
        "blocks": blocks,
        "quality_rows": dataset["quality_rows"],
        "warnings": [
            "14个连续数据块不是14种已确认处理条件；按源表顺序编码为G01-G14，名称为数据块01-14，metadata_status=unknown。",
            "4行长度分布占比合计偏离100%超过1个百分点；原值保留，进入质量复核视图。",
            "三项纸张数据是抗张强度、耐破强度、撕裂强度，不与原核心表的三项指数混用。",
        ],
    }


def write_report(report: dict[str, Any], report_dir: Path) -> tuple[Path, Path]:
    report_dir.mkdir(parents=True, exist_ok=True)
    json_path = report_dir / "recycled_summary_precheck_report.json"
    md_path = report_dir / "recycled_summary_precheck_report.md"
    json_path.write_text(
        json.dumps(report, ensure_ascii=False, indent=2, default=_json_value) + "\n",
        encoding="utf-8",
    )

    counts = report["counts"]
    lines = [
        "# 数据汇总工作簿入库预检报告",
        "",
        f"- 状态：**{report['status']}**",
        f"- 来源文件：`{report['source']['file_name']}` / `{report['source']['sheet_name']}`",
        f"- SHA-256：`{report['source']['file_sha256']}`",
        "",
        "## 数量",
        "",
        "| 项目 | 数量 |",
        "|---|---:|",
        f"| 数据集 | {counts['datasets']} |",
        f"| 数据块 | {counts['blocks']} |",
        f"| 原始观测 | {counts['observations']} |",
        f"| 数据块×循环汇总点 | {counts['summary_points']} |",
        f"| 质量复核行 | {counts['quality_review_rows']} |",
        "",
        "## 数据块映射",
        "",
        "| 编号 | 显示名称 | 源表行 | 区段 | 原料 | 循环数 | 每循环重复数 | 条件状态 |",
        "|---|---|---:|---:|---|---:|---:|---|",
    ]
    for block in report["blocks"]:
        lines.append(
            f"| {block['block_code']} | {block['block_name']} "
            f"| {block['source_start_row']}-{block['source_end_row']} "
            f"| {block['section_no']} | {block['raw_material_grade'] or '待确认'} "
            f"| {block['cycle_count']} | {block['expected_replicates_per_cycle']} "
            f"| {block['metadata_status']} |"
        )

    lines.extend(
        [
            "",
            "## 质量复核行",
            "",
            "| Excel行 | 数据块 | 循环 | 重复 | 长度分布合计/% | 相对100%的偏差 |",
            "|---:|---|---:|---:|---:|---:|",
        ]
    )
    for row in report["quality_rows"]:
        lines.append(
            f"| {row['source_row']} | {row['block_code']} | {row['cycle_no']} "
            f"| {row['replicate_no']} | {row['distribution_sum_pct']} | {row['deviation_pct']} |"
        )
    lines.extend(["", "## 告警", ""])
    lines.extend(f"- {warning}" for warning in report["warnings"])
    lines.extend(
        [
            "",
            "> 预检没有修改Excel或数据库。正式写入必须显式执行导入程序的 `--apply` 模式。",
            "",
        ]
    )
    md_path.write_text("\n".join(lines), encoding="utf-8")
    return json_path, md_path


def _sql_literal(value: Any) -> str:
    if value is None:
        return "NULL"
    if isinstance(value, Decimal):
        return format(value, "f")
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return str(value)
    text = str(value).replace("'", "''")
    return f"'{text}'"


def _values(rows: Iterable[Iterable[Any]]) -> str:
    return ",\n".join(
        "    (" + ", ".join(_sql_literal(value) for value in row) + ")"
        for row in rows
    )


def generate_import_sql(dataset: dict[str, Any]) -> str:
    dataset_row = dataset["dataset"]
    blocks = dataset["blocks"]
    observations = dataset["observations"]

    dataset_columns = (
        "dataset_code",
        "study_name",
        "file_name",
        "sheet_name",
        "source_version",
        "file_sha256",
        "workbook_creator",
        "workbook_last_modified_by",
        "notes",
    )
    block_columns = (
        "dataset_id",
        "block_code",
        "block_name",
        "block_order",
        "section_no",
        "contributor_name",
        "raw_material_grade",
        "material_assignment_basis",
        "condition_label",
        "metadata_status",
        "cycle_count",
        "expected_replicates_per_cycle",
        "source_start_row",
        "source_end_row",
        "notes",
    )
    observation_columns = (
        "dataset_id",
        "block_id",
        "cycle_no",
        "replicate_no",
        "source_row",
        *VALUE_COLUMNS,
    )
    dataset_id = (
        "(SELECT dataset_id FROM fiber.recycled_dataset "
        f"WHERE dataset_code = {_sql_literal(DATASET_CODE)})"
    )

    block_rows = []
    for row in blocks:
        block_rows.append(
            (
                _Sql(dataset_id),
                row["block_code"],
                row["block_name"],
                row["block_order"],
                row["section_no"],
                row["contributor_name"],
                row["raw_material_grade"],
                row["material_assignment_basis"],
                row["condition_label"],
                row["metadata_status"],
                row["cycle_count"],
                row["expected_replicates_per_cycle"],
                row["source_start_row"],
                row["source_end_row"],
                row["notes"],
            )
        )

    observation_rows = []
    for row in observations:
        block_id = (
            "(SELECT block_id FROM fiber.recycled_data_block "
            f"WHERE dataset_id = {dataset_id} "
            f"AND block_code = {_sql_literal(row['block_code'])})"
        )
        observation_rows.append(
            (
                _Sql(dataset_id),
                _Sql(block_id),
                row["cycle_no"],
                row["replicate_no"],
                row["source_row"],
                *(row[column] for column in VALUE_COLUMNS),
            )
        )

    def format_row(row: Iterable[Any]) -> str:
        values = []
        for value in row:
            if isinstance(value, _Sql):
                values.append(value.text)
            else:
                values.append(_sql_literal(value))
        return "    (" + ", ".join(values) + ")"

    sql = [
        "\\encoding UTF8",
        "-- 由 recycled_summary_core.py 生成；来源数据逐值可审计。",
        "\\set ON_ERROR_STOP on",
        "BEGIN;",
        "SET LOCAL ROLE fiber_importer;",
        "SET LOCAL search_path TO fiber, public;",
        "",
        "DO $guard$",
        "BEGIN",
        "    IF current_database() <> 'plant_fiber_db' THEN",
        "        RAISE EXCEPTION '只能导入plant_fiber_db';",
        "    END IF;",
        "    IF NOT EXISTS (",
        "        SELECT 1 FROM information_schema.tables",
        "        WHERE table_schema = 'fiber'",
        "          AND table_name = 'recycled_cycle_observation'",
        "    ) THEN",
        "        RAISE EXCEPTION '请先执行V008和V009';",
        "    END IF;",
        "    IF EXISTS (SELECT 1 FROM fiber.recycled_dataset)",
        "       OR EXISTS (SELECT 1 FROM fiber.recycled_data_block)",
        "       OR EXISTS (SELECT 1 FROM fiber.recycled_cycle_observation) THEN",
        "        RAISE EXCEPTION '再生纸浆导入目标表必须为空，拒绝覆盖或重复导入';",
        "    END IF;",
        "END",
        "$guard$;",
        "",
        f"INSERT INTO fiber.recycled_dataset ({', '.join(dataset_columns)})",
        "VALUES",
        "    (" + ", ".join(_sql_literal(dataset_row[column]) for column in dataset_columns) + ");",
        "",
        f"INSERT INTO fiber.recycled_data_block ({', '.join(block_columns)})",
        "VALUES",
        ",\n".join(format_row(row) for row in block_rows) + ";",
        "",
        f"INSERT INTO fiber.recycled_cycle_observation ({', '.join(observation_columns)})",
        "VALUES",
        ",\n".join(format_row(row) for row in observation_rows) + ";",
        "",
        "DO $postcheck$",
        "BEGIN",
        "    IF (SELECT COUNT(*) FROM fiber.recycled_dataset) <> 1",
        "       OR (SELECT COUNT(*) FROM fiber.recycled_data_block) <> 14",
        "       OR (SELECT COUNT(*) FROM fiber.recycled_cycle_observation) <> 131 THEN",
        "        RAISE EXCEPTION '导入后数量检查失败，事务将回滚';",
        "    END IF;",
        "END",
        "$postcheck$;",
        "",
        "COMMIT;",
        "",
    ]
    return "\n".join(sql)


@dataclass(frozen=True)
class _Sql:
    text: str
