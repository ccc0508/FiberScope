from __future__ import annotations

import hashlib
import json
from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


SHEET_NAME = "Sheet1"
SOURCE_VERSION_ORIGINAL = "原始版"
SOURCE_VERSION_BLEND = "0110"
AGGREGATION_METHOD = "剔除偏差值后取最终均值"

FIBER_HEADERS = (
    "纤维长度/mm（算数平均）",
    "纤维长度/mm（长度加权长）",
    "宽度/μm",
    "粗度（mg/m）",
    "扭结角（°）",
    "扭结纤维（%）",
    "弯曲（%）",
    "分丝帚化长度比（%）",
    "折断纤维比（%）",
    "细小纤维（长度比%）",
    "细小纤维（长度加权比%）",
    "细小纤维面积比（%）",
)

PAPER_HEADERS = (
    "定量(g/m2)",
    "厚度(μm)",
    "耐破指数(kPa·m2·g-1)",
    "撕裂指数(mN·m2·g-1)",
    "抗张指数(N·m·g-1)",
    "零距抗张强度(N/cm)",
    "表面粗糙度(ml/min)",
    "透气度",
    "耐折度",
)

FIBER_FIELDS = (
    ("fiber_length_arithmetic_mm", Decimal("0"), None),
    ("fiber_length_weighted_mm", Decimal("0"), None),
    ("fiber_width_um", Decimal("0"), None),
    ("fiber_coarseness_mg_m", Decimal("0"), None),
    ("kink_angle_deg", Decimal("0"), Decimal("180")),
    ("kinked_fiber_pct", Decimal("0"), Decimal("100")),
    ("curved_fiber_pct", Decimal("0"), Decimal("100")),
    ("fibrillation_length_pct", Decimal("0"), Decimal("100")),
    ("broken_fiber_pct", Decimal("0"), Decimal("100")),
    ("fines_length_pct", Decimal("0"), Decimal("100")),
    ("fines_weighted_length_pct", Decimal("0"), Decimal("100")),
    ("fines_area_pct", Decimal("0"), Decimal("100")),
)

PAPER_FIELDS = (
    ("basis_weight_g_m2", Decimal("0"), None),
    ("thickness_um", Decimal("0"), None),
    ("burst_index", Decimal("0"), None),
    ("tear_index", Decimal("0"), None),
    ("tensile_index", Decimal("0"), None),
    ("zero_span_tensile", Decimal("0"), None),
    ("surface_roughness", Decimal("0"), None),
    ("air_permeability", Decimal("0"), None),
    ("folding_endurance_raw", Decimal("0"), None),
)

EXPECTED_COUNTS = {
    "sample": 61,
    "experiment_sample": 63,
    "sample_component": 70,
    "fiber_measurement": 61,
    "paper_measurement": 61,
    "data_source": 178,
}

EXPECTED_BLEND_CODES = (
    "zk010",
    "zk19",
    "zk28",
    "zk37",
    "zk46",
    "zk55",
    "zk64",
    "zk73",
    "zk82",
    "zk91",
    "zk100",
)


@dataclass(frozen=True)
class SourceSpec:
    filename: str
    version: str
    row_count: int
    headers: tuple[Any, ...]
    kind: str


SOURCE_SPECS = (
    SourceSpec(
        "针叶木-打浆转数-打浆度(1).xlsx",
        SOURCE_VERSION_ORIGINAL,
        28,
        ("序号", "打浆转数", "打浆度"),
        "sw_degree",
    ),
    SourceSpec(
        "针叶木-打浆转数-纤维特性(1).xlsx",
        SOURCE_VERSION_ORIGINAL,
        28,
        ("序号", "打浆转数", *FIBER_HEADERS),
        "sw_fiber",
    ),
    SourceSpec(
        "针叶木-打浆转数-纸张性能(1).xlsx",
        SOURCE_VERSION_ORIGINAL,
        28,
        ("序号", "打浆转数/r", *PAPER_HEADERS),
        "sw_paper",
    ),
    SourceSpec(
        "data1-阔叶木-打浆转数-打浆度(1).xlsx",
        SOURCE_VERSION_ORIGINAL,
        24,
        ("序号", "打浆转数/r", "打浆度/°SR"),
        "hw_degree",
    ),
    SourceSpec(
        "data2-阔叶木-打浆转数-纤维特性(1).xlsx",
        SOURCE_VERSION_ORIGINAL,
        24,
        ("序号", "打浆转数/r", *FIBER_HEADERS),
        "hw_fiber",
    ),
    SourceSpec(
        "data3-阔叶木-打浆转数-纸张性能(1).xlsx",
        SOURCE_VERSION_ORIGINAL,
        24,
        ("序号", "打浆转数/r", *PAPER_HEADERS),
        "hw_paper",
    ),
    SourceSpec(
        "data4-针阔-复配-纤维特性0110(1).xlsx",
        SOURCE_VERSION_BLEND,
        11,
        (None, *FIBER_HEADERS),
        "blend_fiber",
    ),
    SourceSpec(
        "data5-针阔-复配-纸张特性0110(1).xlsx",
        SOURCE_VERSION_BLEND,
        11,
        (None, *PAPER_HEADERS),
        "blend_paper",
    ),
)


class ValidationError(RuntimeError):
    pass


@dataclass
class LoadedSource:
    spec: SourceSpec
    path: Path
    rows: list[tuple[int, tuple[Any, ...]]]
    sha256: str


def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and value.strip() == "")


def _is_missing_measurement(value: Any) -> bool:
    return _is_blank(value) or (
        isinstance(value, str) and value.strip() in {"-", "—", "/", "／"}
    )


def _hash_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _decimal(
    value: Any,
    context: str,
    *,
    minimum: Decimal | None = None,
    maximum: Decimal | None = None,
    allow_missing: bool = False,
    warnings: list[str] | None = None,
) -> Decimal | None:
    if _is_missing_measurement(value):
        if allow_missing:
            if warnings is not None:
                warnings.append(f"{context}为空，按规则转换为NULL")
            return None
        raise ValidationError(f"{context}不能为空")

    if isinstance(value, bool) or not isinstance(value, (int, float, Decimal)):
        raise ValidationError(f"{context}必须是Excel数值单元格，当前值为{value!r}")

    try:
        number = Decimal(str(value))
    except InvalidOperation as exc:
        raise ValidationError(f"{context}不是有效数值：{value!r}") from exc

    if not number.is_finite():
        raise ValidationError(f"{context}必须是有限数值")
    if minimum is not None and number < minimum:
        raise ValidationError(f"{context}不能小于{minimum}，当前为{number}")
    if maximum is not None and number > maximum:
        raise ValidationError(f"{context}不能大于{maximum}，当前为{number}")
    return number


def _integer(value: Any, context: str, *, minimum: int = 0) -> int:
    number = _decimal(value, context, minimum=Decimal(minimum))
    assert number is not None
    if number != number.to_integral_value():
        raise ValidationError(f"{context}必须是整数，当前为{number}")
    return int(number)


def _load_source(data_dir: Path, spec: SourceSpec) -> LoadedSource:
    path = data_dir / spec.filename
    if not path.is_file():
        raise ValidationError(f"缺少正式来源文件：{path}")

    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        if SHEET_NAME not in workbook.sheetnames:
            raise ValidationError(f"{spec.filename}缺少有效工作表{SHEET_NAME}")

        for worksheet in workbook.worksheets:
            if worksheet.title == SHEET_NAME:
                continue
            if any(
                not _is_blank(cell.value)
                for row in worksheet.iter_rows()
                for cell in row
            ):
                raise ValidationError(
                    f"{spec.filename}的非正式工作表{worksheet.title}含有数据；"
                    f"请确认后再导入"
                )

        worksheet = workbook[SHEET_NAME]
        all_rows = list(worksheet.iter_rows())
        if not all_rows:
            raise ValidationError(f"{spec.filename}的{SHEET_NAME}为空")

        width = len(spec.headers)
        header_cells = all_rows[0]
        header_values = tuple(
            header_cells[index].value if index < len(header_cells) else None
            for index in range(width)
        )
        if header_values != spec.headers:
            raise ValidationError(
                f"{spec.filename}表头不匹配。实际为{header_values!r}，"
                f"预期为{spec.headers!r}"
            )
        if any(
            not _is_blank(cell.value)
            for cell in header_cells[width:]
        ):
            raise ValidationError(f"{spec.filename}表头存在未映射的额外列")

        result_rows: list[tuple[int, tuple[Any, ...]]] = []
        for row_cells in all_rows[1:]:
            if any(cell.data_type == "f" for cell in row_cells):
                raise ValidationError(
                    f"{spec.filename}第{row_cells[0].row}行含公式；正式来源必须保存静态值"
                )
            values = tuple(
                row_cells[index].value if index < len(row_cells) else None
                for index in range(width)
            )
            extra_values = tuple(cell.value for cell in row_cells[width:])
            if all(_is_blank(value) for value in (*values, *extra_values)):
                continue
            if any(not _is_blank(value) for value in extra_values):
                raise ValidationError(
                    f"{spec.filename}第{row_cells[0].row}行存在未映射的额外列数据"
                )
            result_rows.append((row_cells[0].row, values))

        if len(result_rows) != spec.row_count:
            raise ValidationError(
                f"{spec.filename}应有{spec.row_count}条数据，实际为{len(result_rows)}条"
            )
        expected_excel_rows = list(range(2, spec.row_count + 2))
        actual_excel_rows = [row_number for row_number, _ in result_rows]
        if actual_excel_rows != expected_excel_rows:
            raise ValidationError(
                f"{spec.filename}的数据行必须连续位于第2至{spec.row_count + 1}行"
            )

        return LoadedSource(
            spec=spec,
            path=path,
            rows=result_rows,
            sha256=_hash_file(path),
        )
    finally:
        workbook.close()


def _source_key(source: LoadedSource, excel_row: int) -> tuple[str, str, int, str]:
    return (
        source.spec.filename,
        SHEET_NAME,
        excel_row,
        source.spec.version,
    )


def _measurement_values(
    values: tuple[Any, ...],
    field_specs: tuple[tuple[str, Decimal | None, Decimal | None], ...],
    start_column: int,
    context: str,
    warnings: list[str],
) -> dict[str, Decimal | None]:
    parsed: dict[str, Decimal | None] = {}
    for offset, (field_name, minimum, maximum) in enumerate(field_specs):
        parsed[field_name] = _decimal(
            values[start_column + offset],
            f"{context}字段{field_name}",
            minimum=minimum,
            maximum=maximum,
            allow_missing=True,
            warnings=warnings,
        )
    return parsed


def _assert_unique(rows: Iterable[dict[str, Any]], keys: tuple[str, ...], label: str) -> None:
    seen: set[tuple[Any, ...]] = set()
    for row in rows:
        key = tuple(row[name] for name in keys)
        if key in seen:
            raise ValidationError(f"{label}存在重复键：{key}")
        seen.add(key)


def validate_and_build(data_dir: Path) -> dict[str, Any]:
    data_dir = data_dir.resolve()
    if not data_dir.is_dir():
        raise ValidationError(f"数据目录不存在：{data_dir}")

    loaded = {spec.kind: _load_source(data_dir, spec) for spec in SOURCE_SPECS}
    warnings: list[str] = []
    samples: list[dict[str, Any]] = []
    experiment_samples: list[dict[str, Any]] = []
    components: list[dict[str, Any]] = []
    fiber_measurements: list[dict[str, Any]] = []
    paper_measurements: list[dict[str, Any]] = []
    sources: list[dict[str, Any]] = []

    for spec in SOURCE_SPECS:
        source = loaded[spec.kind]
        for excel_row, _ in source.rows:
            sources.append(
                {
                    "file_name": spec.filename,
                    "sheet_name": SHEET_NAME,
                    "source_row": excel_row,
                    "source_version": spec.version,
                }
            )

    pure_measurements: dict[str, dict[str, dict[str, Decimal | None]]] = {
        "fiber": {},
        "paper": {},
    }

    group_specs = (
        {
            "prefix": "SW",
            "sample_type": "softwood",
            "display_material": "针叶木浆",
            "experiment_code": "SW_BEATING",
            "material_code": "SW",
            "degree": loaded["sw_degree"],
            "fiber": loaded["sw_fiber"],
            "paper": loaded["sw_paper"],
        },
        {
            "prefix": "HW",
            "sample_type": "hardwood",
            "display_material": "阔叶木浆",
            "experiment_code": "HW_BEATING",
            "material_code": "HW",
            "degree": loaded["hw_degree"],
            "fiber": loaded["hw_fiber"],
            "paper": loaded["hw_paper"],
        },
    )

    for group in group_specs:
        degree_source: LoadedSource = group["degree"]
        authoritative: list[tuple[int, int, int, Decimal, str]] = []
        for order, (excel_row, values) in enumerate(degree_source.rows, start=1):
            serial = _integer(
                values[0], f"{degree_source.spec.filename}第{excel_row}行序号", minimum=1
            )
            if serial != order:
                raise ValidationError(
                    f"{degree_source.spec.filename}第{excel_row}行序号应为{order}，实际为{serial}"
                )
            revolutions = _integer(
                values[1], f"{degree_source.spec.filename}第{excel_row}行打浆转数"
            )
            degree = _decimal(
                values[2],
                f"{degree_source.spec.filename}第{excel_row}行打浆度",
                minimum=Decimal("0"),
                maximum=Decimal("100"),
            )
            assert degree is not None
            canonical_code = f"{group['prefix']}_R{revolutions:05d}"
            authoritative.append((serial, excel_row, revolutions, degree, canonical_code))
            samples.append(
                {
                    "canonical_code": canonical_code,
                    "sample_type": group["sample_type"],
                    "display_name": f"{group['display_material']}-{revolutions}r",
                    "notes": None,
                }
            )
            experiment_samples.append(
                {
                    "experiment_code": group["experiment_code"],
                    "canonical_code": canonical_code,
                    "display_code": str(serial),
                    "display_order": serial,
                    "source_key": _source_key(degree_source, excel_row),
                }
            )
            components.append(
                {
                    "canonical_code": canonical_code,
                    "material_code": group["material_code"],
                    "proportion_pct": Decimal("100"),
                    "beating_revolutions": revolutions,
                    "beating_degree_sr": degree,
                    "source_key": _source_key(degree_source, excel_row),
                }
            )

        authoritative_revolutions = [row[2] for row in authoritative]
        if len(set(authoritative_revolutions)) != len(authoritative_revolutions):
            raise ValidationError(f"{degree_source.spec.filename}存在重复打浆转数")

        for measurement_type, source_key_name, field_specs, destination in (
            ("纤维", "fiber", FIBER_FIELDS, fiber_measurements),
            ("纸张", "paper", PAPER_FIELDS, paper_measurements),
        ):
            measurement_source: LoadedSource = group[source_key_name]
            measurement_revolutions: list[int] = []
            for order, (excel_row, values) in enumerate(measurement_source.rows, start=1):
                serial = _integer(
                    values[0],
                    f"{measurement_source.spec.filename}第{excel_row}行序号",
                    minimum=1,
                )
                if serial != order:
                    raise ValidationError(
                        f"{measurement_source.spec.filename}第{excel_row}行序号应为{order}，"
                        f"实际为{serial}"
                    )
                revolutions = _integer(
                    values[1],
                    f"{measurement_source.spec.filename}第{excel_row}行打浆转数",
                )
                measurement_revolutions.append(revolutions)
                canonical_code = f"{group['prefix']}_R{revolutions:05d}"
                measurements = _measurement_values(
                    values,
                    field_specs,
                    2,
                    f"{measurement_source.spec.filename}第{excel_row}行",
                    warnings,
                )
                record: dict[str, Any] = {
                    "canonical_code": canonical_code,
                    **measurements,
                    "source_key": _source_key(measurement_source, excel_row),
                }
                if source_key_name == "paper":
                    record["replicate_count"] = 3
                    record["aggregation_method"] = AGGREGATION_METHOD
                destination.append(record)
                pure_measurements[source_key_name][canonical_code] = measurements

            if measurement_revolutions != authoritative_revolutions:
                raise ValidationError(
                    f"{measurement_source.spec.filename}的打浆转数或顺序与"
                    f"{degree_source.spec.filename}不一致"
                )

    blend_fiber_source = loaded["blend_fiber"]
    blend_paper_source = loaded["blend_paper"]
    actual_fiber_codes = [str(values[0]).strip() for _, values in blend_fiber_source.rows]
    actual_paper_codes = [str(values[0]).strip() for _, values in blend_paper_source.rows]
    if tuple(actual_fiber_codes) != EXPECTED_BLEND_CODES:
        raise ValidationError(
            f"{blend_fiber_source.spec.filename}复配编号或顺序不符合预期："
            f"{actual_fiber_codes}"
        )
    if tuple(actual_paper_codes) != EXPECTED_BLEND_CODES:
        raise ValidationError(
            f"{blend_paper_source.spec.filename}复配编号或顺序不符合预期："
            f"{actual_paper_codes}"
        )

    blend_fiber_by_code: dict[str, dict[str, Decimal | None]] = {}
    blend_paper_by_code: dict[str, dict[str, Decimal | None]] = {}
    for source, field_specs, destination_map in (
        (blend_fiber_source, FIBER_FIELDS, blend_fiber_by_code),
        (blend_paper_source, PAPER_FIELDS, blend_paper_by_code),
    ):
        for excel_row, values in source.rows:
            code = str(values[0]).strip()
            destination_map[code] = _measurement_values(
                values,
                field_specs,
                1,
                f"{source.spec.filename}第{excel_row}行",
                warnings,
            )

    endpoint_links = {
        "zk010": "HW_R08000",
        "zk100": "SW_R14000",
    }
    for blend_code, canonical_code in endpoint_links.items():
        if blend_fiber_by_code[blend_code] != pure_measurements["fiber"][canonical_code]:
            raise ValidationError(
                f"复配端点{blend_code}的纤维数据与{canonical_code}不一致"
            )
        if blend_paper_by_code[blend_code] != pure_measurements["paper"][canonical_code]:
            raise ValidationError(
                f"复配端点{blend_code}的纸张数据与{canonical_code}不一致"
            )

    for order, code in enumerate(EXPECTED_BLEND_CODES, start=1):
        sw_pct = (order - 1) * 10
        hw_pct = 100 - sw_pct
        fiber_excel_row = blend_fiber_source.rows[order - 1][0]
        paper_excel_row = blend_paper_source.rows[order - 1][0]
        if order == 1:
            canonical_code = "HW_R08000"
        elif order == 11:
            canonical_code = "SW_R14000"
        else:
            canonical_code = f"BLEND_SW{sw_pct:03d}_HW{hw_pct:03d}"
            samples.append(
                {
                    "canonical_code": canonical_code,
                    "sample_type": "blend",
                    "display_name": f"针叶木浆{sw_pct}%+阔叶木浆{hw_pct}%",
                    "notes": None,
                }
            )
            for material_code, proportion, revolutions, degree in (
                ("SW", sw_pct, 14000, Decimal("41.0")),
                ("HW", hw_pct, 8000, Decimal("40.6")),
            ):
                components.append(
                    {
                        "canonical_code": canonical_code,
                        "material_code": material_code,
                        "proportion_pct": Decimal(proportion),
                        "beating_revolutions": revolutions,
                        "beating_degree_sr": degree,
                        "source_key": _source_key(blend_fiber_source, fiber_excel_row),
                    }
                )
            fiber_measurements.append(
                {
                    "canonical_code": canonical_code,
                    **blend_fiber_by_code[code],
                    "source_key": _source_key(blend_fiber_source, fiber_excel_row),
                }
            )
            paper_measurements.append(
                {
                    "canonical_code": canonical_code,
                    **blend_paper_by_code[code],
                    "replicate_count": 3,
                    "aggregation_method": AGGREGATION_METHOD,
                    "source_key": _source_key(blend_paper_source, paper_excel_row),
                }
            )

        experiment_samples.append(
            {
                "experiment_code": "BLEND",
                "canonical_code": canonical_code,
                "display_code": code,
                "display_order": order,
                "source_key": _source_key(blend_fiber_source, fiber_excel_row),
            }
        )

    dataset = {
        "sources": sources,
        "samples": samples,
        "experiment_samples": experiment_samples,
        "components": components,
        "fiber_measurements": fiber_measurements,
        "paper_measurements": paper_measurements,
        "loaded_sources": loaded,
        "warnings": warnings,
    }

    _assert_unique(sources, ("file_name", "sheet_name", "source_row", "source_version"), "data_source")
    _assert_unique(samples, ("canonical_code",), "sample")
    _assert_unique(
        experiment_samples,
        ("experiment_code", "canonical_code"),
        "experiment_sample",
    )
    _assert_unique(components, ("canonical_code", "material_code"), "sample_component")
    _assert_unique(fiber_measurements, ("canonical_code",), "fiber_measurement")
    _assert_unique(paper_measurements, ("canonical_code",), "paper_measurement")

    actual_counts = {
        "sample": len(samples),
        "experiment_sample": len(experiment_samples),
        "sample_component": len(components),
        "fiber_measurement": len(fiber_measurements),
        "paper_measurement": len(paper_measurements),
        "data_source": len(sources),
    }
    if actual_counts != EXPECTED_COUNTS:
        raise ValidationError(
            f"标准化记录数量不符合设计。实际：{actual_counts}，预期：{EXPECTED_COUNTS}"
        )

    ratios: dict[str, Decimal] = defaultdict(lambda: Decimal("0"))
    for component in components:
        ratios[component["canonical_code"]] += component["proportion_pct"]
    bad_ratios = {code: total for code, total in ratios.items() if total != Decimal("100")}
    if bad_ratios:
        raise ValidationError(f"存在组分比例合计不等于100%的样本：{bad_ratios}")

    return dataset


def build_report(dataset: dict[str, Any], data_dir: Path) -> dict[str, Any]:
    loaded: dict[str, LoadedSource] = dataset["loaded_sources"]
    counts = {
        "sample": len(dataset["samples"]),
        "experiment_sample": len(dataset["experiment_samples"]),
        "sample_component": len(dataset["components"]),
        "fiber_measurement": len(dataset["fiber_measurements"]),
        "paper_measurement": len(dataset["paper_measurements"]),
        "data_source": len(dataset["sources"]),
    }
    known_warnings = [
        "R016：透气度源表未标单位，数值原样保留，设计暂按μm/(Pa·s)记录",
        "R017：耐折度单位待确认，数值原样写入folding_endurance_raw",
    ]
    warnings = [*dataset["warnings"], *known_warnings]
    checks = [
        {"rule_id": f"R{number:03d}", "status": "PASS"}
        for number in range(1, 16)
    ]
    checks.extend(
        [
            {"rule_id": "R016", "status": "WARNING"},
            {"rule_id": "R017", "status": "WARNING"},
            {"rule_id": "R018", "status": "PASS"},
            {"rule_id": "R019", "status": "PASS"},
            {"rule_id": "R020", "status": "PASS"},
            {"rule_id": "R021", "status": "NOT_CHECKED", "reason": "由数据库权限脚本V004检查"},
            {"rule_id": "R022", "status": "NOT_CHECKED", "reason": "应在备份交付阶段执行"},
        ]
    )
    return {
        "status": "PASS",
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "data_directory": str(data_dir.resolve()),
        "blocking_error_count": 0,
        "warning_count": len(warnings),
        "warnings": warnings,
        "counts": {
            name: {
                "expected": EXPECTED_COUNTS[name],
                "actual": counts[name],
                "matched": counts[name] == EXPECTED_COUNTS[name],
            }
            for name in EXPECTED_COUNTS
        },
        "source_files": [
            {
                "file_name": spec.filename,
                "version": spec.version,
                "sheet_name": SHEET_NAME,
                "expected_rows": spec.row_count,
                "actual_rows": len(loaded[spec.kind].rows),
                "sha256": loaded[spec.kind].sha256,
            }
            for spec in SOURCE_SPECS
        ],
        "checks": checks,
    }


def write_report(report: dict[str, Any], output_dir: Path) -> tuple[Path, Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    json_path = output_dir / "precheck_report.json"
    markdown_path = output_dir / "precheck_report.md"
    json_path.write_text(
        json.dumps(report, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    lines = [
        "# 植物纤维数据入库预检报告",
        "",
        f"- 状态：{report['status']}",
        f"- 生成时间（UTC）：{report['generated_at_utc']}",
        f"- 数据目录：`{report['data_directory']}`",
        f"- 阻断错误：{report['blocking_error_count']}",
        f"- 告警：{report['warning_count']}",
        "",
        "## 数量核对",
        "",
        "| 对象 | 预期 | 实际 | 结果 |",
        "|---|---:|---:|---|",
    ]
    for name, item in report["counts"].items():
        lines.append(
            f"| {name} | {item['expected']} | {item['actual']} | "
            f"{'通过' if item['matched'] else '失败'} |"
        )
    lines.extend(["", "## 来源文件", "", "| 文件 | 版本 | 数据行 | SHA-256 |", "|---|---|---:|---|"])
    for source in report["source_files"]:
        lines.append(
            f"| {source['file_name']} | {source['version']} | {source['actual_rows']} | "
            f"`{source['sha256']}` |"
        )
    lines.extend(["", "## 告警", ""])
    lines.extend(f"- {warning}" for warning in report["warnings"])
    lines.extend(
        [
            "",
            "## 结论",
            "",
            "8份正式来源通过只读预检，可以进入事务化正式入库。"
            "透气度和耐折度的单位状态不阻断数值入库。",
            "",
        ]
    )
    markdown_path.write_text("\n".join(lines), encoding="utf-8")
    return json_path, markdown_path


@dataclass(frozen=True)
class SqlExpression:
    text: str


def _sql_value(value: Any) -> str:
    if isinstance(value, SqlExpression):
        return value.text
    if value is None:
        return "NULL"
    if isinstance(value, Decimal):
        return format(value, "f")
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, str):
        return "'" + value.replace("'", "''") + "'"
    raise TypeError(f"不支持的SQL值类型：{type(value)!r}")


def _insert_statement(table: str, columns: tuple[str, ...], rows: list[tuple[Any, ...]]) -> str:
    rendered_rows = [
        "    (" + ", ".join(_sql_value(value) for value in row) + ")"
        for row in rows
    ]
    return (
        f"INSERT INTO fiber.{table} (\n    "
        + ",\n    ".join(columns)
        + "\n)\nVALUES\n"
        + ",\n".join(rendered_rows)
        + ";\n"
    )


def _sample_id(code: str) -> SqlExpression:
    return SqlExpression(
        "(SELECT sample_id FROM fiber.sample WHERE canonical_code = "
        f"{_sql_value(code)})"
    )


def _experiment_id(code: str) -> SqlExpression:
    return SqlExpression(
        "(SELECT experiment_id FROM fiber.experiment WHERE experiment_code = "
        f"{_sql_value(code)})"
    )


def _material_id(code: str) -> SqlExpression:
    return SqlExpression(
        "(SELECT material_id FROM fiber.material WHERE material_code = "
        f"{_sql_value(code)})"
    )


def _source_id(key: tuple[str, str, int, str]) -> SqlExpression:
    file_name, sheet_name, source_row, source_version = key
    return SqlExpression(
        "(SELECT source_id FROM fiber.data_source WHERE file_name = "
        f"{_sql_value(file_name)} AND sheet_name = {_sql_value(sheet_name)} "
        f"AND source_row = {source_row} AND source_version = {_sql_value(source_version)})"
    )


def generate_import_sql(dataset: dict[str, Any]) -> str:
    statements = [
        r"\encoding UTF8",
        r"\set ON_ERROR_STOP 1",
        "",
        "-- Generated only after all eight official Excel files passed validation.",
        "-- Run with the fiber_importer login. The six business tables must be empty.",
        "BEGIN;",
        "SET LOCAL search_path TO fiber, public;",
        "SET CONSTRAINTS ALL DEFERRED;",
        "",
        "DO $preflight$",
        "BEGIN",
        "    IF current_database() <> 'plant_fiber_db' THEN",
        "        RAISE EXCEPTION 'Import must run in plant_fiber_db';",
        "    END IF;",
        "    IF current_user <> 'fiber_importer' THEN",
        "        RAISE EXCEPTION 'Import must run as fiber_importer; current user is %', current_user;",
        "    END IF;",
        "    IF (SELECT COUNT(*) FROM fiber.material) <> 2",
        "       OR (SELECT COUNT(*) FROM fiber.experiment) <> 3 THEN",
        "        RAISE EXCEPTION 'Reference data is incomplete; run V003 first';",
        "    END IF;",
        "    IF (SELECT COUNT(*) FROM fiber.sample) <> 0",
        "       OR (SELECT COUNT(*) FROM fiber.data_source) <> 0",
        "       OR (SELECT COUNT(*) FROM fiber.experiment_sample) <> 0",
        "       OR (SELECT COUNT(*) FROM fiber.sample_component) <> 0",
        "       OR (SELECT COUNT(*) FROM fiber.fiber_measurement) <> 0",
        "       OR (SELECT COUNT(*) FROM fiber.paper_measurement) <> 0 THEN",
        "        RAISE EXCEPTION 'Business tables are not empty; import stopped to prevent overwrite';",
        "    END IF;",
        "END",
        "$preflight$;",
        "",
    ]

    statements.append(
        _insert_statement(
            "data_source",
            ("file_name", "sheet_name", "source_row", "source_version"),
            [
                (
                    row["file_name"],
                    row["sheet_name"],
                    row["source_row"],
                    row["source_version"],
                )
                for row in dataset["sources"]
            ],
        )
    )
    statements.append(
        _insert_statement(
            "sample",
            ("canonical_code", "sample_type", "display_name", "notes"),
            [
                (
                    row["canonical_code"],
                    row["sample_type"],
                    row["display_name"],
                    row["notes"],
                )
                for row in dataset["samples"]
            ],
        )
    )
    statements.append(
        _insert_statement(
            "experiment_sample",
            ("experiment_id", "sample_id", "display_code", "display_order", "source_id"),
            [
                (
                    _experiment_id(row["experiment_code"]),
                    _sample_id(row["canonical_code"]),
                    row["display_code"],
                    row["display_order"],
                    _source_id(row["source_key"]),
                )
                for row in dataset["experiment_samples"]
            ],
        )
    )
    statements.append(
        _insert_statement(
            "sample_component",
            (
                "sample_id",
                "material_id",
                "proportion_pct",
                "beating_revolutions",
                "beating_degree_sr",
                "source_id",
            ),
            [
                (
                    _sample_id(row["canonical_code"]),
                    _material_id(row["material_code"]),
                    row["proportion_pct"],
                    row["beating_revolutions"],
                    row["beating_degree_sr"],
                    _source_id(row["source_key"]),
                )
                for row in dataset["components"]
            ],
        )
    )

    fiber_columns = tuple(field[0] for field in FIBER_FIELDS)
    statements.append(
        _insert_statement(
            "fiber_measurement",
            ("sample_id", *fiber_columns, "source_id"),
            [
                (
                    _sample_id(row["canonical_code"]),
                    *(row[column] for column in fiber_columns),
                    _source_id(row["source_key"]),
                )
                for row in dataset["fiber_measurements"]
            ],
        )
    )

    paper_columns = tuple(field[0] for field in PAPER_FIELDS)
    statements.append(
        _insert_statement(
            "paper_measurement",
            (
                "sample_id",
                *paper_columns,
                "replicate_count",
                "aggregation_method",
                "source_id",
            ),
            [
                (
                    _sample_id(row["canonical_code"]),
                    *(row[column] for column in paper_columns),
                    row["replicate_count"],
                    row["aggregation_method"],
                    _source_id(row["source_key"]),
                )
                for row in dataset["paper_measurements"]
            ],
        )
    )

    statements.extend(
        [
            "SET CONSTRAINTS ALL IMMEDIATE;",
            "",
            "DO $count_check$",
            "BEGIN",
            "    IF (SELECT COUNT(*) FROM fiber.sample) <> 61",
            "       OR (SELECT COUNT(*) FROM fiber.experiment_sample) <> 63",
            "       OR (SELECT COUNT(*) FROM fiber.sample_component) <> 70",
            "       OR (SELECT COUNT(*) FROM fiber.fiber_measurement) <> 61",
            "       OR (SELECT COUNT(*) FROM fiber.paper_measurement) <> 61",
            "       OR (SELECT COUNT(*) FROM fiber.data_source) <> 178 THEN",
            "        RAISE EXCEPTION 'Imported row counts do not match the approved baseline';",
            "    END IF;",
            "END",
            "$count_check$;",
            "",
            "COMMIT;",
            "",
            "SELECT 'IMPORT_PASS' AS import_status,",
            "       (SELECT COUNT(*) FROM fiber.sample) AS sample_count,",
            "       (SELECT COUNT(*) FROM fiber.experiment_sample) AS experiment_sample_count,",
            "       (SELECT COUNT(*) FROM fiber.sample_component) AS component_count,",
            "       (SELECT COUNT(*) FROM fiber.fiber_measurement) AS fiber_count,",
            "       (SELECT COUNT(*) FROM fiber.paper_measurement) AS paper_count,",
            "       (SELECT COUNT(*) FROM fiber.data_source) AS source_count;",
            "",
        ]
    )
    return "\n".join(statements)


def write_failure_report(error: Exception, data_dir: Path, output_dir: Path) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / "precheck_report.json"
    payload = {
        "status": "FAIL",
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "data_directory": str(data_dir.resolve()),
        "blocking_error_count": 1,
        "errors": [str(error)],
    }
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return path
